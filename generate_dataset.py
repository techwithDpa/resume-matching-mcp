"""Generate a synthetic resume + job-description dataset for the RAG matcher.

Outputs:
  resumes/<name>.(txt|docx|pdf)   100 resumes, section-structured, mixed formats
  jobs/<id>.txt                   8 job descriptions
  dataset_meta.json               ground-truth metadata for evaluation

Run:
  python generate_dataset.py

Requires python-docx (DOCX) and fpdf2 (PDF).
"""

from __future__ import annotations

import json
import os
import random

import docx
from fpdf import FPDF
from fpdf.enums import XPos, YPos

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
RESUME_DIR = os.path.join(BASE_DIR, "resumes")
JOB_DIR = os.path.join(BASE_DIR, "jobs")
META_PATH = os.path.join(BASE_DIR, "dataset_meta.json")

# role -> (title, core skills, extra skills, education field)
ROLES = {
    "backend": ("Backend Engineer",
                ["Python", "Django", "PostgreSQL", "REST APIs", "Redis"],
                ["Celery", "gRPC", "Microservices", "Docker"], "Computer Science"),
    "ml": ("Machine Learning Engineer",
           ["Python", "Machine Learning", "PyTorch", "scikit-learn", "NLP"],
           ["TensorFlow", "MLOps", "Pandas", "Spark"], "Machine Learning"),
    "data_science": ("Data Scientist",
                     ["Python", "Statistics", "Pandas", "SQL", "Machine Learning"],
                     ["R", "A/B Testing", "Tableau", "Experimentation"], "Statistics"),
    "frontend": ("Frontend Engineer",
                 ["JavaScript", "React", "TypeScript", "CSS", "HTML"],
                 ["Redux", "Next.js", "Accessibility", "Webpack"], "Computer Science"),
    "fullstack": ("Full Stack Developer",
                  ["JavaScript", "React", "Node.js", "TypeScript", "MongoDB"],
                  ["GraphQL", "Express", "AWS", "Docker"], "Software Engineering"),
    "devops": ("DevOps Engineer",
               ["AWS", "Terraform", "Docker", "Kubernetes", "CI/CD"],
               ["Ansible", "Python", "Go", "Prometheus"], "Computer Engineering"),
    "sre": ("Site Reliability Engineer",
            ["Kubernetes", "Prometheus", "Go", "Python", "Linux"],
            ["Grafana", "Terraform", "Incident Response", "SLOs"], "Computer Science"),
    "data_eng": ("Data Engineer",
                 ["Python", "Spark", "Airflow", "SQL", "ETL"],
                 ["Kafka", "Snowflake", "dbt", "AWS"], "Data Engineering"),
    "security": ("Security Engineer",
                 ["Python", "Threat Modeling", "Penetration Testing", "SIEM", "AWS"],
                 ["Burp Suite", "Go", "Cryptography", "SOC"], "Cybersecurity"),
    "mobile": ("Mobile Engineer",
               ["Swift", "Kotlin", "iOS", "Android", "Mobile"],
               ["React Native", "Flutter", "REST APIs", "CI/CD"], "Computer Science"),
    "qa": ("QA Automation Engineer",
           ["Python", "Selenium", "Test Automation", "CI/CD", "pytest"],
           ["Cypress", "Playwright", "API Testing", "Performance Testing"], "Information Technology"),
    "pm": ("Product Manager",
           ["Roadmapping", "Analytics", "SQL", "A/B Testing", "Agile"],
           ["Figma", "Stakeholder Management", "Go-to-Market", "Prioritization"], "Business Administration"),
}

N_RESUMES = 100
SEED = 42

# years of experience -> (seniority level, title prefix)
YEARS_CHOICES = [1, 2, 3, 3, 4, 5, 5, 6, 7, 8, 9, 10, 12]


def seniority_for(years: int) -> tuple[str, str]:
    if years < 3:
        return "Junior", "Junior"
    if years < 7:
        return "Mid", ""
    return "Senior", "Senior"


# Skills a candidate may have picked up outside their main role.
CROSS_SKILLS = ["Docker", "AWS", "Kubernetes", "SQL", "Python", "TypeScript",
                "GraphQL", "Terraform", "Go", "Kafka", "React", "Node.js"]

ACHIEVEMENTS = [
    "cut p95 latency by {n}%", "reduced infrastructure cost by {n}%",
    "grew test coverage to {n}%", "shipped {n} production releases per month",
    "mentored {n} engineers", "improved model accuracy by {n} points",
    "led a migration serving {n}M users", "reduced incident count by {n}%",
]

FIRST = ["John", "Jane", "Raj", "Maria", "Liam", "Aisha", "Chen", "Sofia",
         "David", "Emma", "Omar", "Nina", "Carlos", "Priya", "Yuki", "Tariq",
         "Lena", "Diego", "Fatima", "Noah", "Grace", "Ivan", "Mei", "Hassan",
         "Zoe", "Arjun", "Clara", "Sven", "Leila", "Marco", "Anya", "Kofi",
         "Hana", "Pedro", "Sara", "Tom"]
LAST = ["Doe", "Smith", "Patel", "Garcia", "O'Connor", "Khan", "Wei", "Rossi",
        "Mueller", "Wilson", "Haddad", "Volkov", "Ruiz", "Nair", "Tanaka", "Aziz",
        "Berg", "Silva", "Noor", "Brown", "Lee", "Petrov", "Lin", "Ali",
        "Park", "Mehta", "Weiss", "Larsson", "Saab", "Conti", "Ivanova", "Mensah",
        "Kim", "Costa", "Cohen", "Fischer"]

COMPANIES = ["Acme Corp", "Beta Inc", "DataWorks", "CloudNine", "ShipFast",
             "VisionAI", "StreamHub", "BrightUI", "SafeGuard", "GrowthLabs",
             "Nodal", "PixelLabs", "CoreSystems", "Quantify", "NimbusIO"]

FMT_CYCLE = ["txt", "txt", "txt", "docx", "pdf"]  # ~60% txt, 20% docx, 20% pdf


def candidate_name(i: int) -> str:
    """Unique deterministic name; the first 36 match the Milestone 2 dataset."""
    first = FIRST[i % len(FIRST)]
    last = LAST[(i + 5 * (i // len(FIRST))) % len(LAST)]
    return f"{first} {last}"


def build_candidates(n: int = N_RESUMES, seed: int = SEED) -> list[dict]:
    """Build n candidates cycling through roles with varied years and skills.

    Variation is deliberate so ranking is non-trivial: some candidates drop a
    core skill (gaps / borderline cases), some pick up cross-role skills.
    """
    rng = random.Random(seed)
    roles = list(ROLES.items())
    cands = []
    for i in range(n):
        role_key, (title, core, extra, edu) = roles[i % len(roles)]
        years = rng.choice(YEARS_CHOICES)
        level, prefix = seniority_for(years)
        skills = list(core)
        if rng.random() < 0.3:  # missing one core skill
            skills.remove(rng.choice(core[1:]))
        n_extra = {"Junior": 1, "Mid": 2, "Senior": 3}[level]
        skills += rng.sample(extra, n_extra)
        if rng.random() < 0.35:
            cross = rng.choice([s for s in CROSS_SKILLS if s not in skills])
            skills.append(cross)
        achievements = [a.format(n=rng.randint(2, 60)) for a in rng.sample(ACHIEVEMENTS, 2)]
        grad_year = 2026 - years - 4
        cands.append({
            "id": f"{role_key}_{level.lower()}_{i:03d}",
            "name": candidate_name(i),
            "role": role_key,
            "title": f"{prefix} {title}".strip(),
            "seniority": level,
            "years": years,
            "skills": skills,
            "achievements": achievements,
            "education": f"B.S. {edu}, {COMPANIES[i % len(COMPANIES)].split()[0]} University, {grad_year}",
            "company": COMPANIES[i % len(COMPANIES)],
            "company2": COMPANIES[(i + 5) % len(COMPANIES)],
            "fmt": FMT_CYCLE[i % len(FMT_CYCLE)],
        })
    return cands


def resume_lines(c: dict) -> dict:
    """Return the structured sections of a resume as a dict of section->text."""
    start = 2026 - max(1, c["years"] // 2)
    first_start = 2026 - c["years"]
    summary = (f"{c['title']} with {c['years']} years of experience. "
               f"Specializes in {', '.join(c['skills'][:3])}.")
    base_title = c["title"].replace("Senior ", "").replace("Junior ", "")
    experience = [
        f"{c['title']}, {c['company']} ({start}-present): "
        f"delivered projects using {', '.join(c['skills'][:3])}; {c['achievements'][0]}.",
    ]
    if c["years"] >= 2:
        experience.append(
            f"{base_title}, {c['company2']} ({first_start}-{start}): built systems with "
            f"{', '.join(c['skills'][3:6] or c['skills'][:2])}; {c['achievements'][1]}.")
    return {
        "Summary": summary,
        "Skills": ", ".join(c["skills"]),
        "Experience": experience,
        "Education": c["education"],
    }


def render_text(c: dict, s: dict) -> str:
    lines = [c["name"], c["title"], f"{c['name'].split()[0].lower()}.{c['name'].split()[1].lower().strip(chr(39))}@example.com",
             "", "SUMMARY", s["Summary"], "", "SKILLS", s["Skills"], "", "EXPERIENCE"]
    lines += [f"- {e}" for e in s["Experience"]]
    lines += ["", "EDUCATION", s["Education"], ""]
    return "\n".join(lines)


def write_txt(path, c, s):
    with open(path, "w", encoding="utf-8") as f:
        f.write(render_text(c, s))


def write_docx(path, c, s):
    d = docx.Document()
    d.add_heading(c["name"], 0)
    d.add_paragraph(c["title"])
    d.add_heading("Summary", 1); d.add_paragraph(s["Summary"])
    d.add_heading("Skills", 1); d.add_paragraph(s["Skills"])
    d.add_heading("Experience", 1)
    for e in s["Experience"]:
        d.add_paragraph(e, style="List Bullet")
    d.add_heading("Education", 1); d.add_paragraph(s["Education"])
    d.save(path)


def write_pdf(path, c, s):
    nl = {"new_x": XPos.LMARGIN, "new_y": YPos.NEXT}
    pdf = FPDF(); pdf.add_page()
    pdf.set_font("Helvetica", "B", 16); pdf.cell(0, 10, c["name"], **nl)
    pdf.set_font("Helvetica", "", 12); pdf.cell(0, 8, c["title"], **nl); pdf.ln(3)

    def section(title, body):
        pdf.set_font("Helvetica", "B", 13); pdf.cell(0, 8, title, **nl)
        pdf.set_font("Helvetica", "", 11)
        for line in body:
            pdf.set_x(pdf.l_margin)
            pdf.multi_cell(pdf.epw, 6, line.encode("latin-1", "replace").decode("latin-1"),
                           new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.ln(2)

    section("Summary", [s["Summary"]])
    section("Skills", [s["Skills"]])
    section("Experience", [f"- {e}" for e in s["Experience"]])
    section("Education", [s["Education"]])
    pdf.output(path)


WRITERS = {"txt": write_txt, "docx": write_docx, "pdf": write_pdf}

# Job descriptions: target_roles + must-have skills + min years drive evaluation.
JOBS = [
    {"id": "jd_ml_engineer", "title": "Senior Machine Learning Engineer",
     "target_roles": ["ml", "data_science"], "must_have": ["Python", "Machine Learning"],
     "min_years": 5,
     "text": ("We are hiring a Senior Machine Learning Engineer to build and deploy "
              "ML models in production. Must have 5+ years of Python and hands-on "
              "machine learning experience (PyTorch or TensorFlow), including NLP. "
              "Experience with MLOps and large-scale data pipelines is a plus.")},
    {"id": "jd_backend_python", "title": "Backend Engineer (Python)",
     "target_roles": ["backend", "data_eng"], "must_have": ["Python", "REST APIs"],
     "min_years": 3,
     "text": ("Backend Engineer to build scalable REST APIs and services. Strong "
              "Python and relational database (PostgreSQL) skills required, 3+ years. "
              "Familiarity with Redis, Docker, and microservices preferred.")},
    {"id": "jd_frontend_react", "title": "Frontend Engineer (React)",
     "target_roles": ["frontend", "fullstack"], "must_have": ["React", "JavaScript"],
     "min_years": 3,
     "text": ("Frontend Engineer with strong React and JavaScript/TypeScript skills, "
              "3+ years. You will build accessible, performant UIs. Experience with "
              "Next.js and design systems is a plus.")},
    {"id": "jd_devops_sre", "title": "DevOps / SRE",
     "target_roles": ["devops", "sre"], "must_have": ["Kubernetes", "Terraform"],
     "min_years": 4,
     "text": ("DevOps/SRE to own our cloud infrastructure and reliability. Must have "
              "Kubernetes, Terraform, and AWS experience, 4+ years. CI/CD, Prometheus, "
              "and incident response experience valued.")},
    {"id": "jd_data_engineer", "title": "Data Engineer",
     "target_roles": ["data_eng"], "must_have": ["Spark", "Airflow", "SQL"],
     "min_years": 4,
     "text": ("Data Engineer to build batch and streaming pipelines. Must have Spark, "
              "Airflow, and SQL, 4+ years. Kafka, Snowflake, and dbt experience a plus.")},
    {"id": "jd_security", "title": "Security Engineer",
     "target_roles": ["security"], "must_have": ["Threat Modeling", "Penetration Testing"],
     "min_years": 4,
     "text": ("Security Engineer to lead application security. Must have threat "
              "modeling and penetration testing experience, 4+ years. Python and "
              "cloud (AWS) security knowledge required.")},
    {"id": "jd_fullstack_react", "title": "Full Stack Engineer (React + Node.js)",
     "target_roles": ["fullstack", "frontend"], "must_have": ["React", "Node.js"],
     "min_years": 3,
     "text": ("Full Stack Engineer to own features end to end. Must have React and "
              "Node.js with TypeScript, 3+ years. GraphQL, MongoDB, and AWS experience "
              "is a plus.")},
    {"id": "jd_platform_engineer", "title": "Senior Platform Engineer",
     "target_roles": ["sre", "devops"], "must_have": ["Kubernetes", "Go"],
     "min_years": 5,
     "text": ("Senior Platform Engineer to build our internal developer platform. Must "
              "have Kubernetes and Go, 5+ years. Terraform, Prometheus, and AWS "
              "experience preferred; on-call and SLO ownership a plus.")},
]


def main():
    os.makedirs(RESUME_DIR, exist_ok=True)
    os.makedirs(JOB_DIR, exist_ok=True)
    candidates = build_candidates()
    meta = {"resumes": {}, "jobs": {}}

    for c in candidates:
        s = resume_lines(c)
        fname = f"{c['name'].replace(' ', '_').replace(chr(39), '')}.{c['fmt']}"
        path = os.path.join(RESUME_DIR, fname)
        WRITERS[c["fmt"]](path, c, s)
        meta["resumes"][fname] = {
            "name": c["name"], "role": c["role"], "title": c["title"],
            "seniority": c["seniority"], "years": c["years"], "skills": c["skills"],
            "achievements": c["achievements"], "education": c["education"],
        }

    for j in JOBS:
        with open(os.path.join(JOB_DIR, f"{j['id']}.txt"), "w", encoding="utf-8") as f:
            f.write(f"{j['title']}\n\n{j['text']}\n")
        meta["jobs"][j["id"]] = j

    with open(META_PATH, "w", encoding="utf-8") as f:
        json.dump(meta, f, indent=2)

    fmt_counts = {}
    for c in candidates:
        fmt_counts[c["fmt"]] = fmt_counts.get(c["fmt"], 0) + 1
    print(f"Generated {len(candidates)} resumes {fmt_counts} and {len(JOBS)} job descriptions.")
    print("Wrote dataset_meta.json")


if __name__ == "__main__":
    main()
